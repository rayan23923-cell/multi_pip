"""Builds the A9.5.6 acceptance files (decks 1-7, photos) for the real-device run.

Usage: make_acceptance_fixtures.py <output directory>
The PDFs in the acceptance set were exported from decks 1-6 and 7a with LibreOffice.
Copies used by the tests live in App/TaskLensTests/Fixtures as acceptance-*.
"""
import io, zipfile, random, sys, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFont
out = sys.argv[1]; os.makedirs(out, exist_ok=True)
def font(sz):
    for p in ["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"]:
        if os.path.exists(p): return ImageFont.truetype(p, sz)
    return ImageFont.load_default()
def img(w,h,text,color):
    im = Image.new("RGB",(w,h),color); d = ImageDraw.Draw(im)
    d.rectangle([20,20,w-20,h-20], outline="white", width=12)
    d.text((w//2,h//2), text, fill="white", font=font(min(w,h)//5), anchor="mm")
    b = io.BytesIO(); im.save(b,"JPEG",quality=85); return b.getvalue()
def deck():
    p = Presentation(); p.slide_width, p.slide_height = Inches(13.333), Inches(7.5); return p
def title_slide(p, t, sub=""):
    s = p.slides.add_slide(p.slide_layouts[1] if sub else p.slide_layouts[5]); s.shapes.title.text = t
    if sub: s.placeholders[1].text = sub
    return s
# 1 simple
p = deck()
for i in range(1,4): title_slide(p, f"Simple {i} of 3", f"Plain text on slide {i}.")
p.save(f"{out}/1-simple-text.pptx")
# 2 multi
p = deck()
for i in range(1,16): title_slide(p, f"Slide {i} of 15", "• point one\n• point two\n• point three")
p.save(f"{out}/2-multi-slide-15.pptx")
# 3 arabic
p = deck()
ar = ["العرض التقديمي العربي","الشريحة الثانية: نص من اليمين إلى اليسار","الشريحة الثالثة: أرقام ١٢٣ و 123","الشريحة الرابعة: English داخل النص العربي","الشريحة الأخيرة"]
for i,t in enumerate(ar,1):
    s = title_slide(p, t, f"هذه هي الشريحة رقم {i} من {len(ar)}.\nسطر ثانٍ للتحقق من الاتجاه.")
    for ph in s.placeholders:
        for para in ph.text_frame.paragraphs:
            pPr = para._p.get_or_add_pPr(); pPr.set("rtl","1"); pPr.set("algn","r")
p.save(f"{out}/3-arabic-rtl.pptx")
# 4 images/shapes
p = deck(); cols=["#1f6feb","#d1242f","#1a7f37","#8250df","#bf8700","#0a3069"]
for i in range(1,7):
    s = title_slide(p, f"Images and shapes {i} of 6")
    s.shapes.add_picture(io.BytesIO(img(1200,800,f"Photo {i}",cols[i-1])), Inches(0.7), Inches(1.6), width=Inches(6))
    for j,shape in enumerate([MSO_SHAPE.OVAL, MSO_SHAPE.RIGHT_ARROW, MSO_SHAPE.STAR_5_POINT]):
        sh = s.shapes.add_shape(shape, Inches(7.4+j*1.9), Inches(3), Inches(1.6), Inches(1.6))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(cols[(i+j)%6][1:])
p.save(f"{out}/4-images-shapes.pptx")
# 5 tables/charts
p = deck()
s = title_slide(p,"Table 1 of 2"); t = s.shapes.add_table(5,4,Inches(1),Inches(1.8),Inches(11),Inches(4)).table
for r in range(5):
    for c in range(4): t.cell(r,c).text = ["Region","Q1","Q2","Q3"][c] if r==0 else (f"Region {r}" if c==0 else str(random.Random(r*10+c).randint(10,99)))
s = title_slide(p,"Table 2 of 2 (Arabic)"); t = s.shapes.add_table(3,3,Inches(1),Inches(1.8),Inches(11),Inches(3)).table
for r,row in enumerate([["المنطقة","الربع الأول","الربع الثاني"],["الرياض","40","55"],["جدة","32","47"]]):
    for c,v in enumerate(row): t.cell(r,c).text=v
for kind,name in [(XL_CHART_TYPE.COLUMN_CLUSTERED,"Column chart"),(XL_CHART_TYPE.LINE_MARKERS,"Line chart"),(XL_CHART_TYPE.PIE,"Pie chart")]:
    s = title_slide(p,name); cd = CategoryChartData(); cd.categories=["Q1","Q2","Q3","Q4"]; cd.add_series("Sales",(19,34,27,42))
    if kind!=XL_CHART_TYPE.PIE: cd.add_series("Costs",(12,18,22,25))
    s.shapes.add_chart(kind,Inches(1.5),Inches(1.6),Inches(10),Inches(5.4),cd)
p.save(f"{out}/5-tables-charts.pptx")
# 6 large
p = deck()
for i in range(1,81):
    s = title_slide(p, f"Large deck slide {i} of 80")
    s.shapes.add_picture(io.BytesIO(img(1600,900,f"{i}",cols[i%6])), Inches(2.5), Inches(1.6), width=Inches(8.3))
p.save(f"{out}/6-large-80-slides.pptx")
# 7 problematic
p = deck()
for i in range(1,5):
    s = title_slide(p, f"Deck with speaker notes {i} of 4", "This deck has speaker notes.")
    s.notes_slide.notes_text_frame.text = f"Speaker note for slide {i}."
p.save(f"{out}/7a-problem-speaker-notes.pptx")
open(f"{out}/7b-problem-not-a-presentation.pptx","wb").write(b"This is plain text renamed to .pptx, not a PowerPoint file.\n"*20)
src = zipfile.ZipFile(f"{out}/1-simple-text.pptx")
with zipfile.ZipFile(f"{out}/7c-problem-macro.pptx","w",zipfile.ZIP_DEFLATED) as z:
    for n in src.namelist(): z.writestr(n, src.read(n))
    z.writestr("ppt/vbaProject.bin", b"\xd0\xcf\x11\xe0 fake macro project")
data = open(f"{out}/1-simple-text.pptx","rb").read(); open(f"{out}/7d-problem-truncated.pptx","wb").write(data[:len(data)//2])
# images
os.makedirs(f"{out}/images",exist_ok=True)
sizes=[(4032,3024),(3024,4032),(1920,1080),(1080,1920),(2000,2000),(4000,1000)]
for i,(w,h) in enumerate(sizes,1):
    open(f"{out}/images/image-{i}-{w}x{h}.jpg","wb").write(img(w,h,f"Image {i}",cols[i-1]))
Image.new("RGB",(1500,1000),"#57606a").save(f"{out}/images/image-7-png.png")
