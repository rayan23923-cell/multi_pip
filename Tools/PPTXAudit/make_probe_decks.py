#!/usr/bin/env python3
"""Builds the extra decks the A8 probe opens: 10, 50 and 100 slides, plus
hostile files (macro extension, garbage, ZIP bomb, XXE).

Usage: make_probe_decks.py <audit deck> <output directory>
"""
import io
import shutil
import sys
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

audit, out = Path(sys.argv[1]), Path(sys.argv[2])
out.mkdir(parents=True, exist_ok=True)


def scale_deck(count):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for n in range(1, count + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1)).text_frame
        box.text = f"Slide {n} — الشريحة {n}"
        box.paragraphs[0].runs[0].font.size = Pt(36)
        if n % 3 == 0:
            # A different 1920×1080 image on every third slide, as a photo-heavy deck would have.
            image = Image.new("RGB", (1920, 1080), ((n * 37) % 256, (n * 91) % 256, (n * 53) % 256))
            draw = ImageDraw.Draw(image)
            for x in range(0, 1920, 24):
                draw.line([(x, 0), (1920 - x, 1080)], fill=(255, 255, 255), width=2)
            data = io.BytesIO()
            image.save(data, "JPEG", quality=85)
            data.seek(0)
            slide.shapes.add_picture(data, Inches(0.5), Inches(1.5), Inches(8))
        else:
            body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5)).text_frame
            body.word_wrap = True
            body.text = ("هذا نص عربي للاختبار. " if n % 2 else "English body text for the test. ") * 12
    path = out / f"scale-{count}.pptx"
    prs.save(path)
    print(path, path.stat().st_size, "bytes")


for count in (10, 50, 100):
    scale_deck(count)

# Same deck under the macro-enabled extension (no macro inside; tests what the
# renderers do with the type).
shutil.copy(audit, out / "macro.pptm")

# Not a ZIP at all.
(out / "garbage.pptx").write_bytes(b"This is not a PowerPoint file. " * 200)


def rewrite(target, change):
    with zipfile.ZipFile(audit) as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as copy:
        for item in source.infolist():
            data = source.read(item.filename)
            copy.writestr(item, change(item.filename, data))
    print(target, target.stat().st_size, "bytes")


# ZIP bomb: slide 1's XML grows to about 300 MB of whitespace once inflated.
def bomb(name, data):
    if name == "ppt/slides/slide1.xml":
        head, tail = data.split(b"<p:cSld", 1)
        return head + b" " * 300_000_000 + b"<p:cSld" + tail
    return data


rewrite(out / "zipbomb.pptx", bomb)


# XXE: an external entity naming a local file, used as slide 1's text.
def xxe(name, data):
    if name == "ppt/slides/slide1.xml":
        decl_end = data.index(b"?>") + 2
        data = (data[:decl_end] + b'\n<!DOCTYPE p:sld [<!ENTITY xxe SYSTEM "file:///etc/hosts">]>'
                + data[decl_end:])
        return data.replace(b"TaskLens PowerPoint audit.", b"XXE[&xxe;]")
    return data


rewrite(out / "xxe.pptx", xxe)
